from __future__ import annotations

import base64
import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from accessibility_qa import (  # noqa: E402
    SCHEMA,
    SCHEMA_VERSION,
    VERSION,
    _is_generic_alt_text,
    audit_presentation,
)


_PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9Z8S8AAAAASUVORK5CYII="
)


def _c_nv_pr(shape):
    for element in shape._element.iter():
        if element.tag.rsplit("}", 1)[-1] == "cNvPr":
            return element
    raise AssertionError("shape has no cNvPr element")


def _set_metadata(shape, *, title: str | None = None, description: str | None = None) -> None:
    c_nv_pr = _c_nv_pr(shape)
    if title is not None:
        c_nv_pr.set("title", title)
    if description is not None:
        c_nv_pr.set("descr", description)


def _clear_metadata(shape) -> None:
    c_nv_pr = _c_nv_pr(shape)
    c_nv_pr.attrib.pop("title", None)
    c_nv_pr.attrib.pop("descr", None)


def _add_text(
    slide,
    text: str,
    *,
    left: float,
    top: float,
    width: float,
    height: float,
    font_pt: float,
    name: str | None = None,
    bold: bool = False,
):
    shape = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.text_frame.clear()
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(font_pt)
    run.font.bold = bold
    if name:
        shape.name = name
    return shape


def _add_title(slide, text: str = "Accessible slide title"):
    return _add_text(
        slide,
        text,
        left=0.5,
        top=0.35,
        width=12.2,
        height=0.65,
        font_pt=28,
        name="Slide Title",
        bold=True,
    )


def _add_table(
    slide,
    *,
    left: float,
    top: float,
    width: float = 5.8,
    height: float = 1.6,
    font_pt: float = 12,
):
    shape = slide.shapes.add_table(
        2,
        2,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    values = (("Region", "Revenue"), ("North", "$10M"))
    for row_index, row in enumerate(shape.table.rows):
        for column_index, cell in enumerate(row.cells):
            cell.text_frame.clear()
            run = cell.text_frame.paragraphs[0].add_run()
            run.text = values[row_index][column_index]
            run.font.name = "Arial"
            run.font.size = Pt(font_pt)
    return shape


def _add_chart(slide, *, left: float, top: float, width: float, height: float):
    data = ChartData()
    data.categories = ["A", "B"]
    data.add_series("Series 1", (1, 2))
    return slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
        data,
    )


def _add_group_diagram(slide, *, left: float, top: float):
    group = slide.shapes.add_group_shape()
    group.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(1.2),
        Inches(0.7),
    )
    group.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left + 1.6),
        Inches(top),
        Inches(1.2),
        Inches(0.7),
    )
    return group


class AccessibilityQATests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.image_path = self.root / "pixel.png"
        self.image_path.write_bytes(_PNG_1X1)

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def _presentation(self) -> Presentation:
        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        return presentation

    def _save_passing_deck(self, path: Path) -> None:
        presentation = self._presentation()

        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        _add_title(slide, "Revenue overview")
        _add_text(
            slide,
            "Performance remains above plan.",
            left=0.7,
            top=1.25,
            width=5.5,
            height=0.5,
            font_pt=14,
        )
        picture = slide.shapes.add_picture(
            str(self.image_path), Inches(8.8), Inches(1.25), Inches(3.4), Inches(2.1)
        )
        _set_metadata(
            picture,
            description="Revenue bars compare the north and south regions for the quarter.",
        )
        _add_text(
            slide,
            "Quarterly revenue by region",
            left=0.8,
            top=3.25,
            width=5.8,
            height=0.4,
            font_pt=14,
            bold=True,
        )
        _add_table(slide, left=0.8, top=3.75)

        metadata_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        _add_title(metadata_slide, "Table metadata")
        table = _add_table(metadata_slide, left=0.8, top=1.5)
        _set_metadata(
            table,
            title="Quarterly revenue table",
            description="Revenue totals by region in millions of dollars.",
        )
        presentation.save(path)

    def _save_missing_semantics_deck(self, path: Path) -> None:
        presentation = self._presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])

        picture = slide.shapes.add_picture(
            str(self.image_path), Inches(0.5), Inches(0.7), Inches(3.2), Inches(2.0)
        )
        _clear_metadata(picture)
        chart = _add_chart(slide, left=4.2, top=0.7, width=3.4, height=2.0)
        _clear_metadata(chart)
        diagram = _add_group_diagram(slide, left=8.2, top=0.9)
        _clear_metadata(diagram)
        presentation.save(path)

    def _save_warning_deck(self, path: Path) -> None:
        presentation = self._presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])

        picture = slide.shapes.add_picture(
            str(self.image_path), Inches(0.4), Inches(1.2), Inches(8.0), Inches(3.0)
        )
        # python-pptx writes the source filename as descr; that is still generic alt text.
        self.assertEqual(_c_nv_pr(picture).get("descr"), self.image_path.name)
        _add_title(slide, "Headline follows the visual")
        _add_table(slide, left=9.0, top=4.2, width=3.4, height=1.7)
        presentation.save(path)

    def test_passing_deck_accepts_descriptive_alt_and_both_table_label_paths(self) -> None:
        path = self.root / "passing.pptx"
        self._save_passing_deck(path)

        report = audit_presentation(path)

        self.assertEqual(report["schema"], SCHEMA)
        self.assertEqual(report["version"], VERSION)
        self.assertEqual(report["schema_version"], SCHEMA_VERSION)
        self.assertTrue(report["passed"])
        self.assertEqual(report["error_count"], 0)
        self.assertEqual(report["warning_count"], 0)
        self.assertEqual(report["findings"], [])

    def test_file_paths_are_not_accepted_as_meaningful_alt_text(self) -> None:
        self.assertTrue(_is_generic_alt_text("figure.png"))
        self.assertTrue(_is_generic_alt_text("/tmp/deck-assets/figure.png"))
        self.assertTrue(_is_generic_alt_text(r"C:\\deck-assets\\figure.png"))
        self.assertFalse(_is_generic_alt_text("Microscopy field showing clustered blue nuclei"))

    def test_missing_title_and_alt_text_fail_for_image_chart_and_diagram(self) -> None:
        path = self.root / "missing.pptx"
        self._save_missing_semantics_deck(path)

        report = audit_presentation(path)
        codes = [finding["code"] for finding in report["findings"]]

        self.assertFalse(report["passed"])
        self.assertEqual(report["error_count"], 4)
        self.assertEqual(codes.count("missing_slide_title"), 1)
        self.assertEqual(codes.count("missing_alt_text"), 3)
        self.assertEqual(
            {
                finding["details"].get("visual_kind")
                for finding in report["findings"]
                if finding["code"] == "missing_alt_text"
            },
            {"image", "chart", "diagram"},
        )

    def test_slide_metadata_and_decorative_allow_signals_suppress_false_positives(self) -> None:
        path = self.root / "exemptions.pptx"
        presentation = self._presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])

        picture = slide.shapes.add_picture(
            str(self.image_path), Inches(0), Inches(0), Inches(13.333), Inches(7.5)
        )
        _clear_metadata(picture)
        picture.name = "decorative-background"
        _add_title(slide, "Title after a decorative background")

        metadata_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        metadata_slide._element.cSld.set("name", "Opening overview")
        chart = _add_chart(metadata_slide, left=0.8, top=1.5, width=4.0, height=2.5)
        _set_metadata(chart, description="decorative")
        diagram = _add_group_diagram(metadata_slide, left=6.0, top=1.7)
        _set_metadata(diagram, title="a11y:allow")
        presentation.save(path)

        report = audit_presentation(path)

        self.assertTrue(report["passed"])
        self.assertEqual(report["findings"], [])

    def test_strict_mode_upgrades_generic_alt_table_and_reading_order_warnings(self) -> None:
        path = self.root / "warnings.pptx"
        self._save_warning_deck(path)

        normal = audit_presentation(path)
        strict = audit_presentation(path, strict=True)
        expected_codes = {
            "generic_alt_text",
            "reading_order_visual_before_headline",
            "table_missing_accessible_context",
        }

        self.assertTrue(normal["passed"])
        self.assertEqual(normal["error_count"], 0)
        self.assertEqual(normal["warning_count"], 3)
        self.assertEqual({item["code"] for item in normal["findings"]}, expected_codes)
        self.assertFalse(strict["passed"])
        self.assertEqual(strict["error_count"], 3)
        self.assertEqual(strict["warning_count"], 0)
        self.assertEqual({item["code"] for item in strict["findings"]}, expected_codes)
        self.assertTrue(all(item["base_severity"] == "warning" for item in strict["findings"]))

    def test_font_thresholds_distinguish_body_metadata_and_decorative_numbering(self) -> None:
        path = self.root / "font-sizes.pptx"
        presentation = self._presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        _add_title(slide)
        _add_text(
            slide,
            "Small body copy",
            left=0.8,
            top=1.5,
            width=4.0,
            height=0.5,
            font_pt=10,
        )
        _add_text(
            slide,
            "Source: internal analysis",
            left=0.8,
            top=6.7,
            width=5.0,
            height=0.3,
            font_pt=7,
            name="Source Footer",
        )
        _add_text(
            slide,
            "07",
            left=12.2,
            top=6.8,
            width=0.5,
            height=0.3,
            font_pt=6,
            name="Slide Number Decorative",
        )
        presentation.save(path)

        default_report = audit_presentation(path)
        configured_report = audit_presentation(path, min_body_pt=10, min_metadata_pt=7)

        size_findings = [
            finding
            for finding in default_report["findings"]
            if finding["code"] == "text_below_minimum"
        ]
        self.assertEqual(len(size_findings), 2)
        self.assertEqual(
            {finding["details"]["role"] for finding in size_findings},
            {"body", "metadata"},
        )
        self.assertNotIn("07", {finding["details"]["text_sample"] for finding in size_findings})
        self.assertTrue(configured_report["passed"])
        self.assertEqual(configured_report["finding_count"], 0)

    def test_cli_exit_codes_json_determinism_and_source_immutability(self) -> None:
        path = self.root / "cli.pptx"
        first_report_path = self.root / "first.json"
        second_report_path = self.root / "second.json"
        self._save_warning_deck(path)
        before_hash = hashlib.sha256(path.read_bytes()).hexdigest()

        base_command = [
            sys.executable,
            str(SCRIPTS / "accessibility_qa.py"),
            "--input",
            str(path),
        ]
        first = subprocess.run(
            [*base_command, "--report", str(first_report_path)],
            cwd=ROOT,
            check=False,
            capture_output=True,
            text=True,
        )
        second = subprocess.run(
            [*base_command, "--report", str(second_report_path)],
            cwd=ROOT,
            check=False,
            capture_output=True,
            text=True,
        )
        strict = subprocess.run(
            [*base_command, "--strict"],
            cwd=ROOT,
            check=False,
            capture_output=True,
            text=True,
        )
        missing = subprocess.run(
            [
                sys.executable,
                str(SCRIPTS / "accessibility_qa.py"),
                "--input",
                str(self.root / "does-not-exist.pptx"),
            ],
            cwd=ROOT,
            check=False,
            capture_output=True,
            text=True,
        )

        self.assertEqual(first.returncode, 0, first.stderr)
        self.assertEqual(second.returncode, 0, second.stderr)
        self.assertEqual(first.stdout, second.stdout)
        self.assertEqual(first_report_path.read_bytes(), second_report_path.read_bytes())
        self.assertEqual(first.stdout, first_report_path.read_text(encoding="utf-8"))
        self.assertEqual(json.loads(first.stdout)["warning_count"], 3)
        self.assertEqual(strict.returncode, 1, strict.stderr)
        self.assertEqual(json.loads(strict.stdout)["error_count"], 3)
        self.assertEqual(missing.returncode, 2, missing.stderr)
        self.assertEqual(json.loads(missing.stdout)["findings"][0]["code"], "audit_failed")

        overwrite = subprocess.run(
            [*base_command, "--report", str(path)],
            cwd=ROOT,
            check=False,
            capture_output=True,
            text=True,
        )
        self.assertEqual(overwrite.returncode, 2, overwrite.stderr)
        self.assertIn(
            "must not overwrite",
            json.loads(overwrite.stdout)["findings"][0]["details"]["reason"],
        )
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), before_hash)
        self.assertEqual(audit_presentation(path), audit_presentation(path))


if __name__ == "__main__":
    unittest.main()
