from __future__ import annotations

import sys
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from visual_review import _analyze_text_shapes  # noqa: E402


def _text(
    slide,
    text: str,
    y: float,
    h: float,
    font_pt: float,
    *,
    x: float = 0.5,
    w: float = 9.0,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(font_pt)
    return shape


class VisualReviewTests(unittest.TestCase):
    def test_footer_clearance_does_not_compare_a_low_register_to_itself(self) -> None:
        presentation = Presentation()
        presentation.slide_width = Inches(10.0)
        presentation.slide_height = Inches(5.625)
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        _text(slide, "Assay readout", 0.2, 0.5, 28)
        _text(slide, "Source: data/assay.csv", 4.405, 0.30, 9)
        _text(slide, "METHOD | CONTROL | n | UNCERTAINTY", 4.865, 0.16, 8)
        _text(slide, "Sources: data/assay.csv", 5.325, 0.30, 8)

        issues = _analyze_text_shapes(presentation)

        self.assertNotIn("footer_clearance_risk", {item["type"] for item in issues})

    def test_shrink_to_fit_pressure_is_information_not_false_clip_warning(self) -> None:
        presentation = Presentation()
        presentation.slide_width = Inches(10.0)
        presentation.slide_height = Inches(5.625)
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        shape = _text(
            slide,
            "A long editable sentence that intentionally relies on PowerPoint shrink to fit.",
            2.0,
            0.30,
            12,
            w=3.0,
        )
        shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

        issues = _analyze_text_shapes(presentation)

        pressure = [item for item in issues if item["type"] == "text_autofit_pressure"]
        self.assertTrue(pressure)
        self.assertTrue(all(item["severity"] == "info" for item in pressure))
        self.assertNotIn("text_box_clip_risk", {item["type"] for item in issues})

    def test_title_clearance_ignores_non_overlapping_side_rail_text(self) -> None:
        presentation = Presentation()
        presentation.slide_width = Inches(10.0)
        presentation.slide_height = Inches(5.625)
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        _text(slide, "Main title", 0.3, 0.7, 28, x=0.5, w=5.0)
        _text(slide, "SIDE RAIL", 0.5, 1.0, 12, x=8.0, w=1.2)

        issues = _analyze_text_shapes(presentation)

        self.assertNotIn("title_clearance_risk", {item["type"] for item in issues})


if __name__ == "__main__":
    unittest.main()
