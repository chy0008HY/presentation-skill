from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from design_rules_qa import (  # noqa: E402
    _text_role,
    check_table_caption_overlap,
    check_table_readability,
)
from finalize_quick_deck import _completion_status  # noqa: E402


class TextRoleTests(unittest.TestCase):
    def _shape(self, *, name: str, top: float = 2.0, height: float = 0.58):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        shape = slide.shapes.add_textbox(
            Inches(1.0), Inches(top), Inches(4.0), Inches(height)
        )
        shape.name = name
        shape.text = "TABLE INDEX"
        return shape

    def test_named_metadata_is_caption_even_in_a_tall_box(self) -> None:
        shape = self._shape(name="metadata:table-index-detail")
        self.assertEqual(_text_role(shape, shape.text, 7.5), "caption")

    def test_ordinary_text_in_the_same_box_is_body(self) -> None:
        shape = self._shape(name="Text 1")
        self.assertEqual(_text_role(shape, shape.text, 7.5), "body")


class TableGeometryTests(unittest.TestCase):
    def test_row_heights_cannot_exceed_table_frame(self) -> None:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        table_shape = slide.shapes.add_table(
            3, 2, Inches(1.0), Inches(1.0), Inches(5.0), Inches(1.0)
        )
        for row in table_shape.table.rows:
            row.height = Inches(0.5)
        table_shape.height = Inches(1.0)

        issues = check_table_readability(0, slide, {"min_caption_pt": 7.5})

        self.assertTrue(
            any(issue["type"] == "table_rows_exceed_frame" for issue in issues)
        )

    def test_named_table_caption_cannot_overlap_table_frame(self) -> None:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_table(
            2, 2, Inches(1.0), Inches(1.0), Inches(5.0), Inches(1.5)
        )
        caption = slide.shapes.add_textbox(
            Inches(1.0), Inches(2.3), Inches(5.0), Inches(0.3)
        )
        caption.name = "metadata:table-caption"
        caption.text = "Source: synthetic test data"
        text_shapes = list(
            (index, shape, shape.text)
            for index, shape in enumerate(slide.shapes, start=1)
            if getattr(shape, "has_text_frame", False) and shape.text
        )

        issues = check_table_caption_overlap(0, slide, text_shapes)

        self.assertTrue(any(issue["type"] == "table_caption_overlap" for issue in issues))


class FinalizerReceiptTests(unittest.TestCase):
    def test_render_only_failure_is_deferred_without_runtime_probing(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            qa_dir = Path(tmp)
            (qa_dir / "qa_report.json").write_text(
                json.dumps({"render_rc": 1, "visual_review_warning_count": 1}),
                encoding="utf-8",
            )
            status = _completion_status(
                [{"stage": "qa", "returncode": 1}], qa_dir
            )
        self.assertEqual(status["failure_category"], "render_environment")
        self.assertEqual(status["render_status"], "deferred_environment")
        self.assertIn("Do not probe", status["next_action"])

    def test_static_qa_findings_remain_source_repair_work(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            qa_dir = Path(tmp)
            (qa_dir / "qa_report.json").write_text(
                json.dumps({"render_rc": 1, "overflow_count": 1}),
                encoding="utf-8",
            )
            status = _completion_status(
                [{"stage": "qa", "returncode": 1}], qa_dir
            )
        self.assertEqual(status["failure_category"], "qa_findings")
        self.assertIn("edit outline.json", status["next_action"])


if __name__ == "__main__":
    unittest.main()
