from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from office_package_hash import office_package_normalized_sha256  # noqa: E402
from reference_deck import inspect_reference_deck, patch_reference_deck  # noqa: E402


class ReferenceDeckTests(unittest.TestCase):
    def _deck(self, path: Path) -> None:
        presentation = Presentation()
        presentation.slide_width = Inches(10)
        presentation.slide_height = Inches(5.625)
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(0.6), Inches(0.5), Inches(8.5), Inches(0.7))
        run = title.text_frame.paragraphs[0].add_run()
        run.text = "Original title"
        run.font.name = "Arial"
        run.font.size = Pt(26)
        run.font.bold = True
        body = slide.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(5.0), Inches(1.1))
        body.text_frame.text = "Untouched body"
        presentation.save(path)

    def test_guarded_text_patch_preserves_geometry_style_and_other_text(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "source.pptx"
            output = root / "output.pptx"
            plan_path = root / "patch.json"
            self._deck(source)
            manifest = inspect_reference_deck(source)
            slide = manifest["slides"][0]
            title = next(element for element in slide["elements"] if element["text"] == "Original title")
            body = next(element for element in slide["elements"] if element["text"] == "Untouched body")
            plan_path.write_text(
                json.dumps(
                    {
                        "schema_version": "reference_deck_patch_v1",
                        "source_manifest_sha256": manifest["manifest_sha256"],
                        "operations": [
                            {
                                "action": "replace_text",
                                "slide_id": slide["slide_id"],
                                "element_id": title["element_id"],
                                "expected_text_sha256": title["text_sha256"],
                                "text": "Updated title",
                            }
                        ],
                    }
                ),
                encoding="utf-8",
            )
            result = patch_reference_deck(input_path=source, patch_path=plan_path, output_path=output)
            self.assertTrue(result["passed"], result["preservation_failures"])
            patched = inspect_reference_deck(output)
            patched_index = {
                element["element_id"]: element
                for element in patched["slides"][0]["elements"]
            }
            self.assertEqual(patched_index[title["element_id"]]["text"], "Updated title")
            self.assertEqual(patched_index[body["element_id"]]["text"], "Untouched body")
            self.assertEqual(patched_index[title["element_id"]]["geometry_sha256"], title["geometry_sha256"])
            self.assertEqual(patched_index[title["element_id"]]["style_sha256"], title["style_sha256"])

    def test_patch_is_reproducible_from_same_source(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "source.pptx"
            first = root / "first.pptx"
            second = root / "second.pptx"
            plan_path = root / "patch.json"
            self._deck(source)
            manifest = inspect_reference_deck(source)
            slide = manifest["slides"][0]
            title = next(element for element in slide["elements"] if element["text"] == "Original title")
            plan_path.write_text(
                json.dumps(
                    {
                        "schema_version": "reference_deck_patch_v1",
                        "operations": [
                            {
                                "action": "replace_text",
                                "slide_id": slide["slide_id"],
                                "element_id": title["element_id"],
                                "expected_text_sha256": title["text_sha256"],
                                "text": "Reproducible title",
                            }
                        ],
                    }
                ),
                encoding="utf-8",
            )
            patch_reference_deck(input_path=source, patch_path=plan_path, output_path=first)
            patch_reference_deck(input_path=source, patch_path=plan_path, output_path=second)
            self.assertEqual(
                office_package_normalized_sha256(first),
                office_package_normalized_sha256(second),
            )

    def test_stale_text_precondition_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "source.pptx"
            plan_path = root / "patch.json"
            self._deck(source)
            manifest = inspect_reference_deck(source)
            slide = manifest["slides"][0]
            title = next(element for element in slide["elements"] if element["text"] == "Original title")
            plan_path.write_text(
                json.dumps(
                    {
                        "schema_version": "reference_deck_patch_v1",
                        "operations": [
                            {
                                "action": "replace_text",
                                "slide_id": slide["slide_id"],
                                "element_id": title["element_id"],
                                "expected_text_sha256": "0" * 64,
                                "text": "Should fail",
                            }
                        ],
                    }
                ),
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "precondition"):
                patch_reference_deck(input_path=source, patch_path=plan_path, output_path=root / "out.pptx")

    def test_reference_workspace_persists_preservation_manifest(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "source.pptx"
            workspace = root / "workspace"
            self._deck(source)

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPTS / "init_deck_workspace.py"),
                    "--workspace",
                    str(workspace),
                    "--title",
                    "Reference edit",
                    "--reference-pptx",
                    str(source),
                ],
                cwd=ROOT,
                check=False,
                capture_output=True,
                text=True,
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            workspace_manifest = json.loads(
                (workspace / "workspace.json").read_text(encoding="utf-8")
            )
            self.assertEqual(
                workspace_manifest["reference_deck_manifest"],
                "reference_deck_manifest.json",
            )
            persisted = json.loads(
                (workspace / "reference_deck_manifest.json").read_text(encoding="utf-8")
            )
            self.assertEqual(
                persisted["manifest_sha256"],
                inspect_reference_deck(source)["manifest_sha256"],
            )


if __name__ == "__main__":
    unittest.main()
