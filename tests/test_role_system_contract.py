from __future__ import annotations

import copy
import json
import subprocess
import sys
import tempfile
import unittest
from collections import Counter
from pathlib import Path

from pptx import Presentation


ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from composition_grammar_catalog import (  # noqa: E402
    V2_ROLE_VARIANT_CANDIDATES,
    build_composition_grammar_catalog,
    quick_deck_agent_brief,
    route_composition_grammars,
    validate_composition_grammar_catalog,
)
from init_deck_workspace import _design_brief_stub, _starter_outline, _style_contract  # noqa: E402
from office_package_hash import office_package_normalized_sha256  # noqa: E402
from style_treatment_profiles import (  # noqa: E402
    PROFILE_OVERRIDES,
    RENDERER_TREATMENT_FIELDS,
    preset_treatment_profile,
)
from taste_grammar_catalog import (  # noqa: E402
    COMPOSITION_GRAMMARS,
    PRESET_TO_GRAMMAR,
    renderer_role_systems_for_preset,
    validate_renderer_role_systems_v1,
    validate_taste_grammar_catalog,
)
from validate_planning import _validate_renderer_role_systems_contract  # noqa: E402
from workflow_atom_context import build_workflow_atom_context, compact_workflow_atom_context  # noqa: E402
from role_layout_contracts import (  # noqa: E402
    ROLE_NAMES as V2_ROLE_NAMES,
    load_role_layout_catalog,
    renderer_role_contracts_for_preset,
    validate_renderer_role_contracts_v2,
)
from preflight import _check_role_variant_alignment  # noqa: E402


EXPECTED_GRAMMARS = {
    "consulting-answer-pyramid",
    "scientific-evidence-plate",
    "clinical-care-pathway",
    "editorial-spread",
    "investor-thesis-stage",
    "operations-grid",
    "policy-public-docket",
    "technical-telemetry-canvas",
}


class TasteGrammarCatalogTests(unittest.TestCase):
    def test_preflight_warns_when_role_bypasses_variant_contract(self) -> None:
        issues = _check_role_variant_alignment(
            {"role": "evidence", "variant": "matrix"},
            3,
        )
        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0]["rule"], "role_variant_contract_mismatch")
        self.assertEqual(issues[0]["severity"], "warning")

        self.assertFalse(
            _check_role_variant_alignment(
                {"role": "decision", "variant": "matrix"},
                3,
            )
        )
        unsupported = _check_role_variant_alignment(
            {"role": "decision", "variant": "kpi-hero"},
            4,
        )
        self.assertEqual(len(unsupported), 1)
        self.assertIn("matrix", unsupported[0]["suggested_fix"])

    def test_policy_table_does_not_invent_a_readout(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            outline = root / "outline.json"
            output = root / "deck.pptx"
            outline.write_text(
                json.dumps(
                    {
                        "title": "Policy table",
                        "deck_style": {
                            "style_preset": "warm-terracotta",
                            "composition_grammar": "policy-public-docket",
                            "readability_contract": {"min_body_pt": 16, "min_metadata_pt": 9},
                        },
                        "slides": [
                            {"type": "title", "role": "title", "title": "Policy table"},
                            {
                                "type": "content",
                                "role": "table",
                                "variant": "table",
                                "title": "Select sites",
                                "headers": ["Site", "Risk", "Call"],
                                "rows": [["Northside", "High", "Launch"], ["Central", "Moderate", "Launch"]],
                            },
                        ],
                    }
                ),
                encoding="utf-8",
            )
            subprocess.run(
                [
                    "node",
                    str(ROOT / "scripts" / "build_deck_pptxgenjs.js"),
                    "--outline",
                    str(outline),
                    "--output",
                    str(output),
                    "--style-preset",
                    "warm-terracotta",
                ],
                cwd=ROOT,
                check=True,
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                text=True,
            )
            slide = Presentation(output).slides[1]
            text = "\n".join(shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False))
            self.assertNotIn("Northside: High", text)

    def test_every_quick_brief_candidate_keeps_its_own_v2_sequence(self) -> None:
        route = route_composition_grammars(
            topic="urban heat resilience",
            user_prompt="public evidence, options, implementation, and sources",
            limit=3,
        )
        brief = quick_deck_agent_brief(route, slide_count=8, agent_profile="quality-first")
        candidates = brief["route_candidates"]
        self.assertGreaterEqual(len(candidates), 2)
        for candidate in candidates:
            self.assertEqual(len(candidate["starter_sequence"]), 8)
            for slide in candidate["starter_sequence"]:
                self.assertIn(slide["variant"], V2_ROLE_VARIANT_CANDIDATES[slide["role"]])
        self.assertNotEqual(
            candidates[0]["starter_sequence"][1]["intent"],
            candidates[1]["starter_sequence"][1]["intent"],
        )

    def test_catalog_cardinality_and_preset_distribution(self) -> None:
        taste_summary = validate_taste_grammar_catalog()
        composition_summary = validate_composition_grammar_catalog()

        self.assertTrue(taste_summary["passed"], taste_summary["failures"])
        self.assertTrue(composition_summary["passed"], composition_summary["failures"])
        self.assertEqual(set(COMPOSITION_GRAMMARS), EXPECTED_GRAMMARS)
        self.assertEqual(taste_summary["grammar_count"], 8)
        self.assertEqual(taste_summary["preset_count"], 13)
        self.assertEqual(set(PRESET_TO_GRAMMAR), set(PROFILE_OVERRIDES))
        self.assertLessEqual(max(Counter(PRESET_TO_GRAMMAR.values()).values()), 2)

        counts = taste_summary["role_system_counts"]
        self.assertEqual(counts["title"], 8)
        self.assertGreaterEqual(counts["section"], 6)
        self.assertGreaterEqual(counts["evidence"], 8)
        self.assertGreaterEqual(counts["data"], 8)
        self.assertGreaterEqual(taste_summary["narrative_arc_count"], 6)

    def test_profiles_add_role_systems_without_removing_legacy_treatments(self) -> None:
        for preset in PROFILE_OVERRIDES:
            with self.subTest(preset=preset):
                profile = preset_treatment_profile(preset)
                self.assertEqual(profile["profile_version"], "deck_preset_treatment_profiles_v1")
                self.assertIn("style_mix_matrix", profile)
                self.assertIn("renderer_treatment_signature", profile)
                self.assertEqual(profile["renderer_treatment_fields"], list(RENDERER_TREATMENT_FIELDS))
                self.assertTrue(
                    set(RENDERER_TREATMENT_FIELDS).issubset(profile["renderer_treatment_defaults"])
                )
                role_systems = profile["renderer_role_systems_v1"]
                self.assertFalse(
                    validate_renderer_role_systems_v1(role_systems, expected_preset=preset)
                )
                role_contracts = profile["renderer_role_contracts_v2"]
                self.assertFalse(
                    validate_renderer_role_contracts_v2(role_contracts, expected_preset=preset)
                )

    def test_routing_honors_preset_lock_and_exposes_role_contract(self) -> None:
        route = route_composition_grammars(
            topic="Assay validation",
            user_prompt="investor market story with a hero metric",
            style_preset="lab-report",
        )
        primary = route["primary"]
        self.assertEqual(primary["grammar_id"], "scientific-evidence-plate")
        self.assertEqual(primary["renderer_role_systems_v1"]["schema_version"], "renderer_role_systems_v1")
        self.assertEqual(primary["renderer_role_contracts_v2"]["schema_version"], "renderer_role_contracts_v2")
        for role in ("title", "section", "evidence", "comparison", "data", "decision", "references"):
            self.assertTrue(primary[f"{role}_system_id"])
        self.assertTrue(primary["narrative_arc"])
        self.assertTrue(primary["density"])
        self.assertTrue(primary["grid"])
        self.assertTrue(primary["reading_path"])
        self.assertTrue(primary["preferred_role_variants"])
        self.assertTrue(primary["invariant_moves"])
        self.assertTrue(primary["forbidden_moves"])

    def test_normal_workflow_persists_role_systems(self) -> None:
        context = compact_workflow_atom_context(
            build_workflow_atom_context(
                user_prompt="Clinical evidence and care pathway review",
                style_preset="executive-clinical",
                include_prompt=False,
            )
        )
        role_systems = context["renderer_role_systems_v1"]
        self.assertEqual(role_systems["composition_grammar_id"], "clinical-care-pathway")
        self.assertEqual(
            context["style_execution_plan"]["renderer_role_systems_v1"],
            role_systems,
        )
        role_contracts = context["renderer_role_contracts_v2"]
        self.assertEqual(role_contracts["composition_grammar_id"], "clinical-care-pathway")
        self.assertEqual(
            context["style_execution_plan"]["renderer_role_contracts_v2"],
            role_contracts,
        )

        brief = _design_brief_stub(
            "Clinical evidence review",
            "executive-clinical",
            user_prompt="Clinical evidence and care pathway review",
        )
        outline = _starter_outline(
            "Clinical evidence review",
            "executive-clinical",
            None,
            None,
            user_prompt="Clinical evidence and care pathway review",
        )
        contract = _style_contract(
            title="Clinical evidence review",
            slug="clinical-evidence-review",
            style_preset="executive-clinical",
            font_pair=None,
            palette_key=None,
            reference_pptx=None,
            user_prompt="Clinical evidence and care pathway review",
        )
        self.assertEqual(brief["style_system"]["renderer_role_systems_v1"], role_systems)
        self.assertEqual(outline["metadata"]["renderer_role_systems_v1"], role_systems)
        self.assertEqual(contract["renderer_role_systems_v1"], role_systems)
        self.assertEqual(brief["style_system"]["renderer_role_contracts_v2"], role_contracts)
        self.assertEqual(outline["metadata"]["renderer_role_contracts_v2"], role_contracts)
        self.assertEqual(contract["renderer_role_contracts_v2"], role_contracts)
        self.assertFalse(_validate_renderer_role_systems_contract(brief))

    def test_auto_topic_routing_keeps_style_and_grammar_dynamic(self) -> None:
        context = compact_workflow_atom_context(
            build_workflow_atom_context(
                user_prompt="Public policy brief on urban heat, equity and budget tradeoffs",
                style_preset="",
                include_prompt=False,
            )
        )
        self.assertEqual(context["target_family"], "forest-research")
        self.assertEqual(
            context["composition_grammar_route"]["primary"]["grammar_id"],
            "policy-public-docket",
        )
        self.assertFalse(context["style_execution_plan"]["explicit_style_lock"])

    def test_auto_workspace_selection_does_not_become_an_explicit_lock(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            workspace = Path(temporary)
            (workspace / "design_brief.json").write_text(
                json.dumps(
                    {
                        "style_system": {
                            "style_preset": "arctic-minimal",
                            "style_selection": {"mode": "auto"},
                        }
                    }
                ),
                encoding="utf-8",
            )
            context = compact_workflow_atom_context(
                build_workflow_atom_context(
                    user_prompt="GPU telemetry latency and incident-response architecture review",
                    workspace=workspace,
                    style_preset="",
                    include_prompt=False,
                )
            )
        self.assertEqual(context["selection_basis"], "workspace_auto_style_preset")
        self.assertEqual(
            context["composition_grammar_route"]["primary"]["grammar_id"],
            "technical-telemetry-canvas",
        )
        self.assertFalse(context["style_execution_plan"]["explicit_style_lock"])

    def test_strict_validation_rejects_unknown_or_inconsistent_system_ids(self) -> None:
        brief = _design_brief_stub("Operations review", "lavender-ops", user_prompt="operations review")
        broken = copy.deepcopy(brief)
        broken["style_system"]["renderer_role_systems_v1"]["title_system_id"] = "title-unknown"
        issues = _validate_renderer_role_systems_contract(broken)
        self.assertTrue(issues)
        self.assertTrue(all(issue["severity"] == "error" for issue in issues))
        self.assertTrue(any("title_system_id" in issue["message"] for issue in issues))

    def test_composition_catalog_records_are_exactly_eight(self) -> None:
        catalog = build_composition_grammar_catalog()
        self.assertEqual(catalog["grammar_count"], 8)
        self.assertEqual(
            {record["grammar_id"] for record in catalog["records"]},
            EXPECTED_GRAMMARS,
        )
        for preset in PROFILE_OVERRIDES:
            self.assertEqual(
                renderer_role_systems_for_preset(preset)["composition_grammar_id"],
                PRESET_TO_GRAMMAR[preset],
            )

    def test_v2_catalog_has_eight_unique_non_overlapping_layouts_per_role(self) -> None:
        catalog = load_role_layout_catalog()
        self.assertEqual(len(catalog["grammars"]), 8)
        for role in V2_ROLE_NAMES:
            systems = {grammar[role]["system_id"] for grammar in catalog["grammars"].values()}
            families = {grammar[role]["layout_family"] for grammar in catalog["grammars"].values()}
            self.assertEqual(len(systems), 8, role)
            self.assertEqual(len(families), 8, role)

    def test_v2_validation_rejects_arbitrary_coordinates(self) -> None:
        payload = renderer_role_contracts_for_preset("lab-report")
        broken = copy.deepcopy(payload)
        broken["roles"]["evidence"]["slots"]["evidence_0"] = [0, 0, 1.4, 1]
        failures = validate_renderer_role_contracts_v2(broken, expected_preset="lab-report")
        self.assertTrue(failures)

    def test_legacy_workspace_upgrade_is_explicit_and_idempotent(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            workspace = Path(temporary)
            v1 = renderer_role_systems_for_preset("lab-report")
            (workspace / "design_brief.json").write_text(
                json.dumps({"style_system": {"style_preset": "lab-report", "renderer_role_systems_v1": v1}}),
                encoding="utf-8",
            )
            (workspace / "style_contract.json").write_text(
                json.dumps({"workspace_version": 1, "build": {"style_preset": "lab-report"}, "renderer_role_systems_v1": v1}),
                encoding="utf-8",
            )
            (workspace / "outline.json").write_text(
                json.dumps({"metadata": {"renderer_role_systems_v1": v1}, "slides": []}),
                encoding="utf-8",
            )
            command = [
                sys.executable,
                str(SCRIPTS / "upgrade_renderer_role_contracts_v2.py"),
                "--workspace",
                str(workspace),
            ]
            first = subprocess.run(command, cwd=ROOT, text=True, capture_output=True, check=False)
            self.assertEqual(first.returncode, 0, first.stdout + first.stderr)
            first_payload = json.loads(first.stdout)
            self.assertEqual(set(first_payload["changed_files"]), {"design_brief.json", "style_contract.json", "outline.json"})
            style_contract = json.loads((workspace / "style_contract.json").read_text(encoding="utf-8"))
            self.assertEqual(style_contract["workspace_version"], 1)
            snapshots = {path.name: path.read_bytes() for path in workspace.glob("*.json")}
            second = subprocess.run(command, cwd=ROOT, text=True, capture_output=True, check=False)
            self.assertEqual(second.returncode, 0, second.stdout + second.stderr)
            second_payload = json.loads(second.stdout)
            self.assertEqual(second_payload["changed_files"], [])
            self.assertEqual(snapshots, {path.name: path.read_bytes() for path in workspace.glob("*.json")})

    def test_v1_only_outline_rebuild_is_normalized_reproducible(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            workspace = Path(temporary)
            v1 = renderer_role_systems_for_preset("lab-report")
            outline = {
                "title": "Archived assay report",
                "metadata": {"renderer_role_systems_v1": v1},
                "slides": [
                    {
                        "type": "title",
                        "role": "title",
                        "title": "Archived assay report",
                        "subtitle": "A v1-only reproducibility fixture.",
                    },
                    {
                        "type": "content",
                        "role": "evidence",
                        "variant": "stats",
                        "title": "The archived result remains stable",
                        "facts": [
                            {"value": "97%", "label": "Agreement", "detail": "frozen fixture"},
                            {"value": "3", "label": "Lots", "detail": "same denominator"},
                            {"value": "0", "label": "Drift", "detail": "normalized package"},
                        ],
                    },
                ],
            }
            outline_path = workspace / "outline.json"
            outline_path.write_text(json.dumps(outline), encoding="utf-8")
            outputs = [workspace / "first.pptx", workspace / "second.pptx"]
            for output in outputs:
                command = [
                    "node",
                    str(SCRIPTS / "build_deck_pptxgenjs.js"),
                    "--outline",
                    str(outline_path),
                    "--output",
                    str(output),
                    "--style-preset",
                    "lab-report",
                ]
                result = subprocess.run(command, cwd=ROOT, text=True, capture_output=True, check=False)
                self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            self.assertEqual(
                office_package_normalized_sha256(outputs[0]),
                office_package_normalized_sha256(outputs[1]),
            )


if __name__ == "__main__":
    unittest.main()
